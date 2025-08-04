from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

STATUS_COLORS = {
    "В наличии": "FFC6EFCE",  # Светло-зеленый
    "Санкционное": "FFFFC7CE",  # Светло-красный
    "Снято с производства": "FFFFE699",  # Светло-желтый
    "Требует запроса": "FFD9D9D9",  # Серый
    "Ошибка анализа": "FFB4C6E7",  # Светло-синий
}

def generate_excel_report(equipment: list, output_path: str):
    """Генерация Excel-отчета с форматированием"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Анализ спецификации"
    
    # Заголовки
    headers = [
        "Наименование",
        "Количество",
        "Ед. изм.",
        "Статус",
        "Цена",
        "Источник",
        "Страница в PDF"
    ]
    ws.append(headers)
    
    # Стили для заголовков
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0070C0", end_color="0070C0", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), 
        right=Side(style="thin"), 
        top=Side(style="thin"), 
        bottom=Side(style="thin")
    )
    
    # Применение стилей к заголовкам
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Добавление данных
    for item in equipment:
        row = [
            item.get("name", ""),
            item.get("quantity", 1),
            item.get("unit", "шт."),
            item.get("status", "Не определен"),
            item.get("price", "N/A"),
            item.get("source", ""),
            item.get("source_page", "")
        ]
        ws.append(row)
    
    # Применение стилей к данным
    for row_idx in range(2, len(equipment) + 2):
        # Цвет по статусу
        status = ws.cell(row=row_idx, column=4).value
        if status in STATUS_COLORS:
            fill = PatternFill(start_color=STATUS_COLORS[status], end_color=STATUS_COLORS[status], fill_type="solid")
            for col in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col).fill = fill
        
        # Границы для всех ячеек
        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).border = thin_border
    
    # Автоширина столбцов
    for col_idx in range(1, len(headers) + 1):
        max_length = 0
        col_letter = get_column_letter(col_idx)
        
        for cell in ws[col_letter]:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        
        adjusted_width = (max_length + 2) * 1.2
        ws.column_dimensions[col_letter].width = adjusted_width
    
    # Фиксация заголовков
    ws.freeze_panes = "A2"
    
    # Сохранение
    wb.save(output_path)